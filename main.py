# Competitive Pricing Analysis System

# Python-Based Competitive Pricing Analytics and Market Positioning Evaluation Framework



# Imports libraries
import shutil
import pandas as pd
import os
import zipfile
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.styles import Alignment
from openpyxl.styles import Font





# Input Data Preparation:

# Gets Excel price files from a ZIP archive and separates the company’s file from competitors’ files.
def get_price_files_from_zip():
    temp_folder = "temp_excel_files"
    while True:
        try:
            zip_path = input("Enter the path to the ZIP file with Excel files: ")
            if os.path.exists(temp_folder):
                shutil.rmtree(temp_folder)
            os.makedirs(temp_folder)
            with zipfile.ZipFile(zip_path, "r") as zip_ref:
                zip_ref.extractall(temp_folder)
            excel_files = []
            for file in os.listdir(temp_folder):
                if file.endswith(".xlsx"):
                    excel_files.append(file)
            if len(excel_files) != 4:
                print("Error: The ZIP file must contain exactly 4 Excel files.")
                print("Available Excel files:", excel_files)
                continue
            print("\nAvailable Excel files:")
            for index, file in enumerate(excel_files, start=1):
                print(f"{index} — {file}")
            our_price_file_index = get_our_price_file_index(excel_files)
            our_price_file = os.path.join(temp_folder, excel_files[our_price_file_index])
            competitor_price_files = []
            for i, file in enumerate(excel_files):
                if i != our_price_file_index:
                    competitor_price_files.append(os.path.join(temp_folder, file))
            return our_price_file, competitor_price_files
        except FileNotFoundError:
            print("File error: The specified ZIP file was not found.")
        except PermissionError:
            print("Permission error: The file cannot be accessed.")
        except OSError:
            print("File system error: Please check the file path and try again.")
        except Exception as e:  # ост случаи
            print(f"Error: Unexpected error - {e}")

# Prompts the user to select the company's price file from the extracted Excel files.
def get_our_price_file_index(excel_files):
    while True:
        try:
            our_price_file_index = int(input("\nEnter the number of your company's price file: ")) - 1
            if our_price_file_index < 0 or our_price_file_index >= len(excel_files):
                print("Error: Please enter a valid file number.")
                continue
        except ValueError:
            print("Error: Please enter a number.")
            continue
        return our_price_file_index

# Obtains the acceptable market price deviation threshold for competitive pricing analysis.
def get_threshold():
    print("\nThreshold explanation:")
    print(
        "\tThe threshold parameter determines the percentage range within which\n"
        "\tthe company’s prices are considered competitively aligned with competitors’ market prices."
    )
    while True:
        threshold_input = input("Enter the threshold as a percentage: ")
        if "," in threshold_input:
            print("Error: Please use a dot as the decimal separator. For example, enter 4.5 instead of 4,5.")
            continue
        try:
            threshold = float(threshold_input)
            if threshold < 0:
                print("Error: The threshold cannot be negative.")
                continue
            if threshold > 50:
                print("Error: The threshold cannot be greater than 50%.")
                continue
            return threshold / 100
        except ValueError:
            print("Error: Please enter a number.")





# Pricing Analysis Engine:

# Validates that all values in the 'price' column are numeric across all input files.
def validate_price_values(our_price_file, competitor_price_files):
    try:
        files = [our_price_file] + competitor_price_files
        for file in files:
            df = pd.read_excel(file)
            invalid_rows = df[df["price"].isna() | pd.to_numeric(df["price"], errors="coerce").isna()]
            if len(invalid_rows) > 0:
                print(f"Validation error in file '{file}': column 'price' must contain numeric values only.")
                return False
        return True
    except Exception as e:
        print(f"Validation error while checking price values: {e}")
        return False

# Merges our price table with competitor price tables into a unified dataset for pricing analysis.
def merge_market_prices(our_price_file, competitor_price_files):
    if not validate_price_values(our_price_file, competitor_price_files):
        return
    df = pd.read_excel(our_price_file)
    competitor_number = 0
    for file in competitor_price_files:
        competitor_number += 1
        df_competitor = pd.read_excel(file)
        df = pd.merge(df, df_competitor, on="item", how="left", suffixes=(None, str(competitor_number)))
    return df.rename(columns={
        "price" : "our_price",
        "price1" : "competitor_price_1",
        "price2" : "competitor_price_2",
        "price3" : "competitor_price_3",
    })

# A helper function used by advanced_price_comparison() to compare our price with a competitor's price.
def basic_price_comparison(our_item_price, competitor_item_price, threshold):
    difference_abs = our_item_price - competitor_item_price
    difference_relative = abs(difference_abs / our_item_price)
    if difference_relative <= threshold:
        return "prices are approximately equal"
    elif our_item_price > competitor_item_price:
        return "our price is higher"
    else:
        return "our price is lower"

# Compares our company’s product price with up to three competitor prices and assigns a pricing category.
def advanced_price_comparison(our_item_price, comp1_item_price, comp2_item_price, comp3_item_price, threshold):
    competitors = [comp1_item_price, comp2_item_price, comp3_item_price]
    available_competitor_prices = []
    for price in competitors:
        if price != 0:
            available_competitor_prices.append(price)
    competitors = available_competitor_prices
    if len(competitors) == 0:
        return "A"
    elif len(competitors) == 1:
        first_comparison_result = basic_price_comparison(our_item_price, competitors[0], threshold)
        if first_comparison_result == "prices are approximately equal":
            return "C"
        elif first_comparison_result == "our price is higher":
            return "D"
        else:
            return "B"
    min_price = min(competitors)
    max_price = max(competitors)
    comparison_with_min = basic_price_comparison(our_item_price, min_price, threshold)
    comparison_with_max = basic_price_comparison(our_item_price, max_price, threshold)
    if comparison_with_min == "prices are approximately equal" and comparison_with_max == "prices are approximately equal":
        return "G"
    elif comparison_with_min == "our price is higher" and comparison_with_max == "our price is lower":
        return "H"
    elif comparison_with_min == "our price is lower":
        return "E"
    elif comparison_with_min == "prices are approximately equal":
        return "F"
    elif comparison_with_max == "our price is higher":
        return "K"
    elif comparison_with_max == "prices are approximately equal":
        return "I"

# Calculates percentage deviation of our price from competitor prices based on the pricing category.
def pricing_deviation(our_item_price, comp1_item_price, comp2_item_price, comp3_item_price, threshold):
    competitors = [comp1_item_price, comp2_item_price, comp3_item_price]
    available_competitor_prices = []
    for price in competitors:
        if price != 0:
            available_competitor_prices.append(price)
    pricing_category = advanced_price_comparison(
        our_item_price,
        comp1_item_price,
        comp2_item_price,
        comp3_item_price,
        threshold
    )
    if pricing_category in ["A"]:
        return "No Competitors"
    elif pricing_category in ["B", "C", "D"]:
        reference_price = available_competitor_prices[0]
    elif pricing_category in ["E", "F"]:
        reference_price = min(available_competitor_prices)
    elif pricing_category in ["I", "K"]:
        reference_price = max(available_competitor_prices)
    elif pricing_category in ["G", "H"]:
        reference_price = sum(available_competitor_prices) / len(available_competitor_prices)
    deviation_percentage = (our_item_price - reference_price) / reference_price * 100
    return str(round(deviation_percentage, 2)) + "%"

# Adds pricing analysis results to the DataFrame:
# 1. pricing_category — a category showing how our price compares with competitors;
# 2. pricing_deviation — percentage deviation from the relevant competitor price benchmark.
def pricing_analysis(df, threshold):
    df = df.fillna(0)  # заполняем пропуски нулями
    df["pricing_category"] = df.apply(
        lambda row: advanced_price_comparison(
            row["our_price"],
            row["competitor_price_1"],
            row["competitor_price_2"],
            row["competitor_price_3"],
            threshold
        ),
        axis=1
    )
    df["pricing_deviation"] = df.apply(
        lambda row: pricing_deviation(
            row["our_price"],
            row["competitor_price_1"],
            row["competitor_price_2"],
            row["competitor_price_3"],
            threshold
        ),
        axis=1
    )
    return df





# Item Category Analysis:

# Calculates the percentage distribution of pricing_category within each item_category.
def pricing_category_distribution_by_item_category(df):
    result = {}
    item_categories = df["item_category"].unique()
    for category in item_categories:
        df_category = df[df["item_category"] == category]
        total_items = len(df_category)
        counts = df_category["pricing_category"].value_counts()
        percentages = ((counts / total_items) * 100).round(2)
        result[category] = percentages.to_dict()
    return result

# Creates an Excel table showing the pricing category distribution table for one item category.
def add_pricing_distribution_table(ws, item_category, percentages, row_position):
    pricing_categories = ["A", "B", "C", "D", "E", "F", "G", "H", "I", "K"]
    color_names = {
        "A": "Red",
        "B": "Green",
        "C": "Green",
        "D": "Green",
        "E": "Blue",
        "F": "Blue",
        "G": "Blue",
        "H": "Blue",
        "I": "Blue",
        "K": "Blue"
    }
    ws.cell(row=row_position, column=1, value=f"Pricing Category Distribution: {item_category}")
    ws.cell(row=row_position + 1, column=1, value="Pricing Category")
    ws.cell(row=row_position + 1, column=2, value="Percentage")
    ws.cell(row=row_position + 1, column=3, value="Color Group")
    for i, cat in enumerate(pricing_categories):
        percent = percentages.get(cat, 0) / 100
        ws.cell(row=row_position + 2 + i, column=1, value=cat)
        cell = ws.cell(row=row_position + 2 + i, column=2, value=percent)
        cell.number_format = "0%"
        cell.alignment = Alignment(horizontal="left")
        ws.cell(row=row_position + 2 + i, column=3, value=color_names[cat])

# Creates an Excel bar chart showing the pricing category distribution for one item category.
def add_pricing_distribution_chart(ws, item_category, row_position):
    pricing_categories = ["A", "B", "C", "D", "E", "F", "G", "H", "I", "K"]
    color_map = {
        "A": "FF0000",
        "B": "00B050",
        "C": "00B050",
        "D": "00B050",
        "E": "0070C0",
        "F": "0070C0",
        "G": "0070C0",
        "H": "0070C0",
        "I": "0070C0",
        "K": "0070C0"
    }
    chart = BarChart()
    chart.width = 16.95
    chart.height = 7.43
    chart.type = "col"
    chart.title = f"Item category: {item_category}"
    chart.title.overlay = False
    chart.legend = None
    chart.y_axis.delete = False
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 1
    chart.y_axis.majorUnit = 0.25
    chart.y_axis.number_format = "0%"
    chart.x_axis.delete = False
    data = Reference(ws, min_col=2, min_row=row_position + 1, max_row=row_position + 11)
    categories = Reference(ws, min_col=1, min_row=row_position + 2, max_row=row_position + 11)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    series = chart.series[0]
    for i, cat in enumerate(pricing_categories):
        point = DataPoint(idx=i)
        point.graphicalProperties.solidFill = color_map[cat]
        series.dPt.append(point)
    ws.add_chart(chart, f"D{row_position}")

# Classifies a percentage value as a strong, moderate, weak, or no signal.
def get_signal_strength(percentage):
    if percentage > 25:
        return "Strong signal"
    elif 10 <= percentage <= 25:
        return "Moderate signal"
    elif 0 < percentage < 10:
        return "Weak signal"
    else:
        return "No signal"

# Returns the recommendation title and text based on the pricing category group and signal strength.
def get_recommendation(group_name, signal):
    if group_name == "A":
        if signal == "Strong signal" or signal == "Moderate signal" or signal == "Weak signal":
            return (
                "Exclusivity pattern:",
                "Products with no directly comparable competitor alternatives are present within the assortment. "
                "The company may leverage this to strengthen product differentiation, "
                "reduce direct price competition, and provide greater pricing flexibility within the category."
            )
        else:
            return (
                "Recommendations are not applicable:",
                "No products were identified in this pricing category group."
            )
    elif group_name == "B+E":
        if signal == "Strong signal" or signal == "Moderate signal" or signal == "Weak signal":
            return (
                "Underpricing pattern:",
                "Products priced below competitor levels are present within the assortment, suggesting potential "
                "unrealized revenue opportunities.The company may consider reviewing selected product prices "
                "to improve margins while maintaining competitiveness."
            )
        else:
            return (
                "Recommendations are not applicable:",
                "No products were identified in this pricing category group."
            )
    elif group_name == "D+K":
        if signal == "Strong signal" or signal == "Moderate signal" or signal == "Weak signal":
            return (
                "Overpricing pattern:",
                "Products priced above competitor levels are present within the assortment, "
                "which may reduce price competitiveness within the category. "
                "The company may consider evaluating selected prices to improve market positioning."
            )
        else:
            return (
                "Recommendations are not applicable:",
                "No products were identified in this pricing category group."
            )
    elif group_name == "C+F+G+H+I":
        if signal == "Strong signal" or signal == "Moderate signal" or signal == "Weak signal":
            return (
                "Market alignment pattern:",
                "Products within market-aligned pricing categories are present within the assortment, "
                "indicating competitive pricing and balanced market positioning. The company may consider "
                "maintaining this approach while continuing to monitor market conditions and competitor pricing."
            )
        else:
            return (
                "Recommendations are not applicable:",
                "No products were identified in this pricing category group."
            )
    return "Something went wrong."

# Adds a recommendation table for one item category based on pricing category group analysis.
def add_recommendations_table(ws, item_category, percentages, row_position):
    start_row = row_position + 15
    groups = {
        "A": percentages.get("A", 0),
        "B+E": percentages.get("B", 0) + percentages.get("E", 0),
        "D+K": percentages.get("D", 0) + percentages.get("K", 0),
        "C+F+G+H+I":
            percentages.get("C", 0)
            + percentages.get("F", 0)
            + percentages.get("G", 0)
            + percentages.get("H", 0)
            + percentages.get("I", 0)
    }
    interpretations = {
        "A": "Exclusive products",
        "B+E": "Underpriced products",
        "D+K": "Overpriced products",
        "C+F+G+H+I": "Market-aligned products"
    }
    headers = ["Pricing Category Group", "Group Interpretation", "Signal Strength", "Recommendation"]
    ws.cell(row=start_row, column=1, value=f"Pricing Recommendations: {item_category}")
    for col, header in enumerate(headers, start=1):
        header_cell = ws.cell(row=start_row + 1, column=col, value=header)
        header_cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[start_row + 1].height = 35
    current_row = start_row + 2
    for group_name, percentage in groups.items():
        signal = get_signal_strength(percentage)
        recommendation_title, recommendation_text = get_recommendation(group_name, signal)
        formatted_text = (recommendation_title + "\n" + recommendation_text)
        ws.cell(row=current_row, column=1, value=group_name)
        ws.cell(row=current_row, column=2, value=interpretations[group_name])
        ws.cell(row=current_row, column=3, value=f"{signal} ({percentage:.1f}%)")
        ws.cell(row=current_row, column=4, value=formatted_text)
        for col in range(1, 5):
            ws.cell(row=current_row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        current_row += 1

# Creates an Excel analysis sheet with pricing distribution tables, charts, and recommendations for each item category.
def create_item_category_analysis_sheet(df, wb):
    percentages_by_category = pricing_category_distribution_by_item_category(df)
    if "Item Category Analysis" in wb.sheetnames:
        del wb["Item Category Analysis"]
    ws = wb.create_sheet("Item Category Analysis")
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 13
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 80
    row_position = 1
    for item_category, percentages in percentages_by_category.items():
        add_pricing_distribution_table(ws, item_category, percentages, row_position)
        add_pricing_distribution_chart(ws, item_category, row_position)
        add_recommendations_table(ws, item_category, percentages, row_position)
        row_position += 26





# Pricing Category Group Reports

# Creates separate DataFrames for each pricing category group.
def create_pricing_group_tables(df):
    pricing_groups = {
        "A": ["A"],
        "B+E": ["B", "E"],
        "D+K": ["D", "K"],
        "C+F+G+H+I": ["C", "F", "G", "H", "I"]
    }
    group_tables = {}
    for group_name, pricing_categories in pricing_groups.items():
        group_tables[group_name] = df[df["pricing_category"].isin(pricing_categories)]
    return group_tables

# Creates Excel worksheets containing products grouped by pricing category.
def create_pricing_group_sheets(df, wb):
    group_tables = create_pricing_group_tables(df)
    for sheet_name, group_df in group_tables.items():
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)
        headers = list(group_df.columns)
        ws.append(headers)
        for row in group_df.values.tolist():
            ws.append(row)





# Pricing Analysis Summary

# Calculates overall pricing category group percentages and summary statistics.
def calculate_pricing_summary_statistics(df):
    pricing_groups = {
        "A": ["A"],
        "B+E": ["B", "E"],
        "D+K": ["D", "K"],
        "C+F+G+H+I": ["C", "F", "G", "H", "I"]
    }
    total_products = len(df)
    group_percentages = {}
    for group_name, pricing_categories in pricing_groups.items():
        group_products = len(df[df["pricing_category"].isin(pricing_categories)])
        percentage = (group_products / total_products) * 100
        group_percentages[group_name] = round(percentage, 2)
    return total_products, group_percentages

# Determines the overall pricing structure based on pricing category group distribution.
def identify_pricing_structure(df):
    total_products, group_percentages = calculate_pricing_summary_statistics(df)
    exclusive = group_percentages["A"]
    underpriced = group_percentages["B+E"]
    overpriced = group_percentages["D+K"]
    market_aligned = group_percentages["C+F+G+H+I"]
    max_value = max(exclusive, underpriced, overpriced, market_aligned)
    highest_group_difference = 5
    highest_groups = []
    values = {
        "exclusive": exclusive,
        "underpriced": underpriced,
        "overpriced": overpriced,
        "market-aligned": market_aligned
    }
    for group_name, percentage in values.items():
        if max_value - percentage <= highest_group_difference:
            highest_groups.append(group_name)
    if market_aligned >= 50 and underpriced <= 10 and overpriced <= 10:
        return "Market-Aligned Pricing Structure"
    if max_value < 35:  # No Clearly Dominant Pricing Category Group
        return "Mixed Pricing Structure"
    if len(highest_groups) == 1:
        if highest_groups[0] == "market-aligned":
            return "Market-Aligned Pricing Structure"
        elif highest_groups[0] == "underpriced":
            return "Underpricing-Oriented Pricing Structure"
        elif highest_groups[0] == "overpriced":
            return "Overpricing-Oriented Pricing Structure"
        elif highest_groups[0] == "exclusive":
            return "Exclusivity-Driven Pricing Structure"
    if "market-aligned" in highest_groups and "underpriced" in highest_groups:
        return "Market-Leading Pricing Structure"
    if "market-aligned" in highest_groups and "overpriced" in highest_groups:
        return "Vulnerable Pricing Structure"
    if "underpriced" in highest_groups and "overpriced" in highest_groups:
        return "Polarized Pricing Structure"
    if "market-aligned" in highest_groups and "exclusive" in highest_groups:
        return "Differentiated Competitive Pricing Structure"
    if "exclusive" in highest_groups and "underpriced" in highest_groups:
        return "Market-Leading Pricing Structure"
    if "exclusive" in highest_groups and "overpriced" in highest_groups:
        return "Out-of-Market Pricing Structure"
    return "Mixed Pricing Structure"

# Identifies potential pricing risks based on the detected pricing structure.
def identify_pricing_risk(df):
    pricing_structure = identify_pricing_structure(df)
    structure_risks = {
        "Market-Aligned Pricing Structure": ["Low Brand Visibility Risk"],

        "Underpricing-Oriented Pricing Structure": ["Low Profitability Risk"],

        "Overpricing-Oriented Pricing Structure": ["Low Demand Risk", "Customer Churn Risk"],

        "Exclusivity-Driven Pricing Structure": ["Low Demand Risk"],

        "Market-Leading Pricing Structure": ["Low Profitability Risk"],

        "Vulnerable Pricing Structure": ["Low Demand Risk", "Customer Churn Risk", "Low Brand Visibility Risk"],

        "Polarized Pricing Structure": ["Low Demand Risk", "Customer Churn Risk", "Low Profitability Risk"],

        "Differentiated Competitive Pricing Structure": ["Low Demand Risk"],

        "Out-of-Market Pricing Structure": ["Low Demand Risk", "Customer Churn Risk"],

        "Mixed Pricing Structure":
            ["Low Demand Risk", "Customer Churn Risk", "Low Profitability Risk", "Low Brand Visibility Risk"]
    }
    return ", ".join(structure_risks[pricing_structure])

# Generates an overall pricing analysis conclusion based on pricing structure and risks.
def get_pricing_analysis_overall(df):
    pricing_balance = identify_pricing_structure(df)
    descriptions = {

        "Market-Aligned Pricing Structure":
            "The assortment remains closely aligned with market pricing levels, resulting in a competitive "
            "position, but potentially limiting the brand’s ability to differentiate itself from competitors.",

        "Underpricing-Oriented Pricing Structure":
            "The assortment contains a substantial share of underpriced products, "
            "which may limit profitability and result in unrealized revenue potential.",

        "Overpricing-Oriented Pricing Structure":
            "The assortment contains a substantial share of overpriced products, "
            "which may reduce demand and increase the risk of customers switching to competitors.",

        "Exclusivity-Driven Pricing Structure":
            "The assortment contains a substantial share of exclusive products with limited direct competition, "
            "which may reduce overall demand due to their niche appeal, but also strengthens product differentiation.",

        "Market-Leading Pricing Structure":
            "The assortment contains a substantial share of products with strong market positions, "
            "though competitive pricing may reduce overall profitability.",

        "Vulnerable Pricing Structure":
            "The assortment contains a substantial share of products with weak competitive positions, "
            "making it more difficult to attract demand, retain customers, and stand out from competitors.",

        "Polarized Pricing Structure":
            "The assortment contains a substantial share of both underpriced and overpriced products, "
            "which may create inconsistent price perceptions and weaken the overall pricing strategy.",

        "Differentiated Competitive Pricing Structure":
            "The assortment contains a substantial share of both market-aligned and exclusive products, "
            "which is generally beneficial, but certain niche products may face limited demand.",

        "Out-of-Market Pricing Structure":
            "The assortment contains a substantial share of both overpriced and exclusive products, "
            "placing it well outside prevailing market expectations and limiting broader customer appeal.",

        "Mixed Pricing Structure":
            "The assortment exhibits a highly mixed pricing structure, "
            "which may confuse customers and limit the effectiveness of the overall pricing strategy."
    }
    return descriptions[pricing_balance]

# Adds a pricing analysis summary table to the worksheet.
def add_pricing_analysis_summary_table(ws, df, row_position):
    start_row = row_position + 18
    total_products, group_percentages = calculate_pricing_summary_statistics(df)
    headers = ["Analysis Component", "Assessment"]
    rows = [
        ["Total products analyzed", total_products],
        ["Exclusive products (A)", group_percentages["A"] / 100],
        ["Underpriced products (B+E)", group_percentages["B+E"] / 100],
        ["Overpriced products (D+K)", group_percentages["D+K"] / 100],
        ["Market-aligned products (C+F+G+H+I)", group_percentages["C+F+G+H+I"] / 100],
        ["Structure", identify_pricing_structure(df)],
        ["Risk", identify_pricing_risk(df)],
        ["Overall Conclusion", get_pricing_analysis_overall(df)]]
    ws.cell(row=start_row, column=1, value="Pricing Analysis Summary")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row + 1, column=col, value=header)
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    for row_index, row_data in enumerate(rows, start=start_row + 2):
        for col_index, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
            if row_data[0] in [
                "Exclusive products (A)",
                "Underpriced products (B+E)",
                "Overpriced products (D+K)",
                "Market-aligned products (C+F+G+H+I)"
            ] and col_index == 2:
                cell.number_format = "0.00%"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 58

# Adds a pricing category group distribution chart to the worksheet.
def add_pricing_analysis_summary_chart(ws, row_position):
    table_start_row = row_position + 18
    short_labels = ["A", "B+E", "D+K", "C+F+G+H+I"]
    for i, label in enumerate(short_labels, start=table_start_row + 3):
        ws.cell(row=i, column=4, value=label)
    data = Reference(ws, min_col=2, min_row=table_start_row + 3, max_row=table_start_row + 6)
    categories = Reference(ws, min_col=4, min_row=table_start_row + 3, max_row=table_start_row + 6)
    chart = BarChart()
    chart.width = 19.1
    chart.height = 9
    chart.type = "col"
    chart.title = "Pricing Category Group Distribution"
    chart.title.overlay = False
    chart.legend = None
    chart.y_axis.delete = False
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 1
    chart.y_axis.majorUnit = 0.25
    chart.y_axis.number_format = "0%"
    chart.x_axis.delete = False
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(categories)
    colors = [
        "FFD966",  # yellow - A
        "5B9BD5",  # blue - B+E
        "FF0000",  # red - D+K
        "70AD47"  # green - C+F+G+H+I
    ]
    series = chart.series[0]
    for i, color in enumerate(colors):
        point = DataPoint(idx=i)
        point.graphicalProperties.solidFill = color
        series.dPt.append(point)
    for cell in ws["D"]:
        cell.font = Font(color="FFFFFF")
    ws.add_chart(chart, f"A{row_position}")

# Creates an Excel worksheet containing the overall pricing analysis summary.
def create_pricing_analysis_summary_sheet(df, wb):
    if "Pricing Analysis Summary" in wb.sheetnames:
        del wb["Pricing Analysis Summary"]
    ws = wb.create_sheet("Pricing Analysis Summary", 0)
    add_pricing_analysis_summary_table(ws, df, 1)
    add_pricing_analysis_summary_chart(ws, 1)





# Report Generation

# Saves the pricing analysis results to an Excel report.
def save_price_analysis_to_excel(df, excel_file_name):
    df.to_excel(excel_file_name, sheet_name="Pricing Analysis Data", index=False, header=True)
    wb = load_workbook(excel_file_name)
    create_item_category_analysis_sheet(df, wb)
    create_pricing_group_sheets(df, wb)
    create_pricing_analysis_summary_sheet(df, wb)
    wb.save(excel_file_name)
    print (f"\nExcel file has been saved as: {excel_file_name}")

# Prompts the user for an output filename and saves the pricing analysis report.
def ask_excel_file_name_and_save_analysis(df):
    print ("\nExample filename:\n\tfinal_price_analysis_report.xlsx")
    while True:
        excel_file_name = input("Enter a filename for the final pricing analysis report: ")
        try:
            save_price_analysis_to_excel(df, excel_file_name)
            break
        except PermissionError as e:
            print(f"Error: Permission denied - {e}")
        except OSError:
            print (f"Error: Invalid filename - {excel_file_name}")
        except Exception as e:
            print(f"Error: Unexpected error - {e}")





# Main Program Flow

our_price_file, competitor_price_files = get_price_files_from_zip()
threshold = get_threshold()
try:
    df = merge_market_prices(our_price_file, competitor_price_files)
    if df is None:
        print ("Analysis stopped due to validation errors")
    else:
        df = pricing_analysis(df, threshold)
        ask_excel_file_name_and_save_analysis(df)
except KeyError:
    print("\nError: Invalid table format in the specified files.")

